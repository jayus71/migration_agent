"""Generate the current repair diagram as editable PowerPoint and vector SVG."""
from pathlib import Path
import xml.etree.ElementTree as ET

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parent
WIDTH, HEIGHT = 12, 5.45
SVG_NS = "http://www.w3.org/2000/svg"
ET.register_namespace("", SVG_NS)


def main():
    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(WIDTH), Inches(HEIGHT)
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    svg = ET.Element(f"{{{SVG_NS}}}svg", width="1200", height="545", viewBox="0 0 1200 545")
    defs = ET.SubElement(svg, "defs")
    marker = ET.SubElement(defs, "marker", id="arrow", markerWidth="8", markerHeight="8",
                           refX="7", refY="4", orient="auto", markerUnits="strokeWidth")
    ET.SubElement(marker, "path", d="M0,0 L8,4 L0,8 Z", fill="#46535F")

    def text(value, x, y, w, h, size=15, bold=False, color="17212B"):
        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        shape.name = value.replace("\n", " | ")
        frame = shape.text_frame
        frame.margin_left = frame.margin_right = 0
        frame.margin_top = frame.margin_bottom = 0
        for i, line in enumerate(value.split("\n")):
            p = frame.paragraphs[0] if i == 0 else frame.add_paragraph()
            p.text = line
            p.font.name = "Times New Roman"
            p.font.size = Pt(size)
            p.font.bold = bold
            p.font.color.rgb = RGBColor.from_string(color)
            p.space_after = Pt(5)
            node = ET.SubElement(svg, "text", x=str(x * 100),
                y=str(y * 100 + size * 100 / 72 + i * (size + 5) * 100 / 72),
                fill="#" + color, **{"font-family": "Times New Roman, serif",
                "font-size": str(size * 100 / 72), "font-weight": "bold" if bold else "normal"})
            node.text = line

    def box(title, lines, x, y, w, fill):
        shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(1.65))
        shape.name = title
        shape.fill.solid()
        shape.fill.fore_color.rgb = RGBColor.from_string(fill)
        shape.line.color.rgb = RGBColor.from_string("A8B2BB")
        shape.line.width = Pt(1)
        ET.SubElement(svg, "rect", x=str(x * 100), y=str(y * 100), width=str(w * 100), height="165",
                      fill="#" + fill, stroke="#A8B2BB", **{"stroke-width": "1.4"})
        text(title, x + .16, y + .12, w - .32, .35, 18, True)
        text(lines, x + .16, y + .65, w - .32, .94, 14)

    def arrow(points):
        for i, (start, end) in enumerate(zip(points, points[1:])):
            connector = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT,
                Inches(start[0]), Inches(start[1]), Inches(end[0]), Inches(end[1]))
            connector.line.color.rgb = RGBColor.from_string("46535F")
            connector.line.width = Pt(1.5)
            if i == len(points) - 2:
                tail = OxmlElement("a:tailEnd")
                tail.set("type", "triangle")
                connector._element.spPr.get_or_add_ln().append(tail)
        ET.SubElement(svg, "polyline", points=" ".join(f"{x*100},{y*100}" for x, y in points),
            fill="none", stroke="#46535F", **{"stroke-width": "2.1", "marker-end": "url(#arrow)"})

    text("LaDiM: autonomous investigation and repair", .16, .10, 11.7, .48, 22, True)
    box("Public inputs", "Source program when available\nInitial candidate and library\nPublic task contract", .15, .8, 3.1, "F1F3F5")
    box("Independent Verifier", "Read and search public code\nRun tests; write scratch tests\nInfer hypotheses and locations", 4, .8, 3.25, "E8F2F9")
    box("Fixer", "Inspect the verifier's evidence\nEdit candidate or library\nRetain history across retries", 8, .8, 3.7, "EDF5EC")
    box("Target execution", "MindSpore via torch4ms\nJAX via TorchAX\nObserve actual backend execution", .15, 3.3, 3.1, "F1F3F5")
    box("External acceptance", "Execution and forward values\nGradients and parameter updates\nContract and confirmation checks", 4, 3.3, 3.25, "FFF2DF")
    box("Workspace and history", "Current candidate and library\nRecorded edits and observations\nShared lifetime budget", 8, 3.3, 3.7, "F1F3F5")
    arrow([(3.25, 1.7), (4, 1.7)])
    arrow([(7.25, 1.7), (8, 1.7)])
    text("evidence", 7.28, 1.22, .72, .3, 11)
    arrow([(9.85, 2.45), (9.85, 3.3)])
    text("syntax-checked edits", 9.99, 2.66, 1.65, .5, 11)
    arrow([(8, 4.35), (7.25, 4.35)])
    arrow([(3.25, 4.1), (4, 4.1)])
    arrow([(7.25, 3.65), (7.61, 3.65), (7.61, 2.05), (8, 2.05)])
    text("failed checks\nand measurements", 5.3, 2.56, 2.25, .6, 12)
    text("Accept after passing checks; stop when the shared budget is exhausted.", 2.1, 5.08, 9.6, .35, 15)
    deck.save(ROOT / "hierarchical_feedback_architecture.pptx")
    ET.indent(svg)
    ET.ElementTree(svg).write(ROOT / "slim_v4_overview.svg", encoding="utf-8", xml_declaration=True)


if __name__ == "__main__":
    main()
